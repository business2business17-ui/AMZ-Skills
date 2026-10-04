import openpyxl,re,collections,json
R="/home/user/business2business17-ui/"
BRAND="MOBILIUS"
FILLER={"and","&","with","of","the","a","an","in","on","for","to"}
_US=[(r"watercolour","watercolor"),(r"colourful","colorful"),(r"colours","colors"),(r"colour","color"),(r"greys","grays"),(r"grey","gray"),(r"centred","centered"),(r"centres","centers"),(r"centre","center"),(r"cosy","cozy")]
def us(t):
    for a,b in _US:
        t=re.sub(a,b,t); t=re.sub(a.capitalize(),b.capitalize(),t)
    return t

# Trigger words removed from print names used in text (title, Item Highlight, bullet 4). Catalog names stay in the source tables.
PRINT_MAP={
 # policy HIGH / MEDIUM / LOW (from owner's 215 policy check)
 "Pink Retro Sticker Collage (Papaya & Strawberry)":("Pink Retro Fruit Collage","policy HIGH: sticker collage"),
 "Matcha Girlie Sticker Collage":("Green Retro Collage","policy HIGH: sticker collage"),
 "Pink Kittens Girly Sticker Collage":("Pink Kittens Collage","policy HIGH: sticker collage"),
 "Retro Tech Tower with Cat (TV, Cassettes, Phones)":("Retro Tech Tower with Cat","policy HIGH: media reference"),
 "Man in Bowler Hat with Green Apple Painting":("Man in Hat with Green Fruit Art","policy HIGH: artwork reference; word Apple"),
 "Japanese JDM Car":("Japanese Sports Car","policy HIGH: car culture term"),
 "Cali Sensor Thermal Landscape":("Thermal Mountain Landscape","policy HIGH: brand-like name"),
 "California National Parks Stickers (Yosemite, Joshua Tree)":("Western Travel Badge Collage","policy HIGH: place names, stickers"),
 "Japanese Retro Sticker Collage (Lucky Cat)":("Japanese Retro Collage","policy MEDIUM: sticker collage"),
 "Teddy Bears & Kittens Sticker Collage":("Teddy Bears & Kittens Collage","policy MEDIUM: sticker collage"),
 "Cherry Bomb Red Sticker Collage":("Red Cherry Collage","policy MEDIUM: song/band-like name"),
 "Surreal Art Sticker Collage (Hearts & Faces)":("Surreal Hearts & Faces Collage","policy MEDIUM: sticker collage"),
 "Funny Dogs Meme Sticker Collage":("Funny Dogs Collage","policy MEDIUM: meme"),
 "Cat Meme Sticker Collage":("Funny Cats Collage","policy MEDIUM: meme"),
 "Anime Girl in Strawberry Shirt":("Girl in Strawberry Shirt Illustration","policy MEDIUM: anime"),
 "Surreal Cat & Guitar Vintage Collage":("Surreal Cat & Music Vintage Collage","policy MEDIUM: guitar trade dress"),
 "Gothic Rabbit & Ghost Girl Woodcut":("Dark Rabbit & Girl Woodcut","policy MEDIUM: horror words"),
 "Crowned Masked Figure Engraving":("Crowned Figure Engraving","policy MEDIUM: mask"),
 "Hands Playing Piano Oil Painting":("Hands Playing Piano Art","policy MEDIUM: painting reference"),
 "Fluffy White Lamb Oil Painting":("Fluffy White Lamb Art","policy MEDIUM: painting reference"),
 "California Travel Stickers (Route 66, Death Valley)":("Retro Travel Stickers Collage","policy MEDIUM: place names"),
 "Chopsticks Holding Mochi Hamster":("Chopsticks Holding Pink Mochi","policy LOW: character"),
 "Reaching Hands Renaissance Fresco":("Reaching Hands Classic Art","policy LOW: artwork reference"),
 "Sunflowers in Vase Oil Painting":("Sunflowers in Vase Art","policy LOW: artwork reference"),
 "Great Japanese Wave Woodblock Print":("Great Japanese Wave Print","policy LOW: artwork reference"),
 "Two Cherub Angels Painting":("Two Cherub Angels Art","policy LOW: artwork reference"),
 "USA Postage Stamps Collage":("Vintage Postage Stamps Collage","policy LOW: postal marks"),
 "Blood Gothic Horror Girl":("Dark Gothic Girl","policy LOW: violent word"),
 "Gothic Bleeding Heart with Thorny Hands":("Dark Gothic Heart with Thorny Hands","policy LOW: violent word"),
 "California Girl Gang License Plate":("Girl Gang License Plate","policy LOW: state plate"),
 "Thin Blue Line Distressed Flag":("Distressed Flag with Blue Stripe","policy LOW: sensitive symbol"),
 "Rejected Movie Script":("Movie Script Page","policy LOW: place name inside text"),
 # 223-only prints not covered by the policy check
 "Japanese Anime Eyes Senpai":("Japanese Comic Eyes","223 only: anime term"),
 "Farmer Couple with Pitchfork Painting":("Farmer Couple with Pitchfork Art","223 only: painting reference"),
 "Acadia National Park Lighthouse & Moose":("Coastal Lighthouse & Moose Scene","223 only: park name"),
 "Grand Canyon National Park with Eagle & Cactus":("Canyon Scene with Eagle & Cactus","223 only: park name"),
 "Joshua Tree National Park Coyote & Roadrunner":("Desert Tree Coyote & Roadrunner Scene","223 only: park name"),
 "Yellowstone National Park Bison & Geyser":("Bison & Geyser Wilderness Scene","223 only: park name"),
 "Yosemite National Park Bear & Waterfall":("Bear & Waterfall Wilderness Scene","223 only: park name"),
 "Zion National Park Bighorn Sheep & Canyon":("Bighorn Sheep & Canyon Scene","223 only: park name"),
 # place names, brand-like words and the word Apple found by name scan
 "Desert Canyon Road (Monument Valley)":("Desert Canyon Road","place name"),
 "El Capitan Topographic Map Carabiner":("Topographic Map Carabiner","place name"),
 "9.0 Richter Vibe California Seismograph":("9.0 Richter Vibe Seismograph","place name"),
 "Stay Cool California Pool Water":("Stay Cool Pool Water","place name"),
 "Stay Cool California Groovy Flowers":("Stay Cool Groovy Flowers","place name"),
 "California Golden State Poppies & Bear":("Golden Poppies & Bear","place name"),
 "California Poppies with Coastal Sunset":("Poppies with Coastal Sunset","place name"),
 "Highway 1 Pacific Coast Highway Badge":("Coastal Road Badge","place name"),
 "HWY 1 Big Sur Coastal Road":("Coastal Cliff Road","place name"),
 "California Sunset Palms":("Sunset Palms","place name"),
 "Vapor Sunset Strip Receipt with Stickers":("Sunset Receipt with Stickers","place name"),
 "Red Apples on Blue Wavy Stripes":("Red Fruit on Blue Wavy Stripes","word Apple"),
 "American Football Rugby Sport Emblem":("Football Rugby Sport Emblem","league-like word"),
 "American Basketball Sport Emblem":("Basketball Sport Emblem","league-like word"),
 "Rugby / American Football Player Sport Graphic":("Rugby Football Player Sport Graphic","league-like word"),
}
def generic_name(pn):
    k=re.sub(r"\s+"," ",pn).strip()
    return PRINT_MAP[k][0] if k in PRINT_MAP else k
def clean_pn(pn):
    return us(re.sub(r"\s+"," ",pn).strip())
def words(pn): return [w for w in re.sub(r"\([^)]*\)","",pn).split()]
def fit_print(pn,budget,taken):
    """full print name if it fits, otherwise a shortened but still unique form"""
    full=re.sub(r"\s+"," ",pn).strip()
    if len(full)<=budget and full.lower() not in taken: return full
    noparen=re.sub(r"\s*\([^)]*\)","",full).strip()
    if len(noparen)<=budget and noparen.lower() not in taken: return noparen
    ow=noparen.split(); cands=[]
    # 1) cut right before a preposition/conjunction (cleanest), longest first
    for k in range(len(ow)-1,0,-1):
        if ow[k].lower() in FILLER and ow[k-1].lower() not in FILLER: cands.append(" ".join(ow[:k]))
    # 2) any prefix that does not end with a filler word
    for k in range(len(ow),0,-1):
        if ow[k-1].lower() not in FILLER: cands.append(" ".join(ow[:k]))
    ws=[w for w in ow if w.lower() not in FILLER] or ow
    for k in range(len(ws)-1,0,-1): cands.append(" ".join([ws[0]]+ws[-k:]))
    for k in range(len(ws)-1,0,-1): cands.append(" ".join(ws[-k:]))
    for c in cands:
        if len(c)<=budget and c.lower() not in taken: return c
    return None
def palette(cn):
    if not cn: return "decorative"
    cn=cn.strip()
    m=re.match(r"Multicolor \((.+)\)$",cn)
    if m:
        p=[x.strip().lower() for x in m.group(1).split("/")]; return "multicolor ("+", ".join(p[:-1])+(" and " if len(p)>1 else "")+p[-1]+")"
    if cn.lower()=="multicolor": return "multicolor"
    mod=""
    m=re.search(r"\((Glitter|Chrome)\)",cn,re.I)
    if m: mod=m.group(1).lower(); cn=re.sub(r"\s*\([^)]*\)","",cn)
    p=[x.strip().lower() for x in cn.split("/") if x.strip()]
    s=(", ".join(p[:-1])+(" and " if len(p)>1 else "")+p[-1]) if p else ""
    if "glitter" in mod or "chrome" in mod: s=(s+" "+mod).strip() if len(p)==1 else s+" with "+mod+" accents"
    return s
MODELS={"Apple iPhone 16e/Apple iPhone 17e":("iPhone 16e/17e","iPhone 16e and iPhone 17e","16e17e"),
 "Apple iPhone 17 Pro / 18 Pro":("iPhone 18/17 Pro","iPhone 18 Pro and iPhone 17 Pro","17 pro"),
 "Apple iPhone 17 Pro Max / 18 Pro Max":("iPhone 18/17 Pro Max","iPhone 18 Pro Max and iPhone 17 Pro Max","17 pro max")}
def minfo(m):
    m=m.strip()
    if m in MODELS: return MODELS[m]
    n=m.replace("Apple ",""); return (n,n,n.replace("iPhone ","").lower())
BK={"13":"funda fundas para forro cover cases iphone13","13 pro max":"funda fundas para forro cover cases","14":"funda fundas para cover cases iphone14","14 pro max":"funda fundas para forro cover cases",
 "15":"funda fundas para forro cover cases iphone15","15 pro":"funda fundas para cover cases","15 pro max":"funda fundas para forro cover cases","16":"funda fundas para forro cover cases iphone16",
 "16 plus":"funda fundas para cover cases","16 pro":"funda fundas para cover cases iphone16pro","16 pro max":"funda fundas para forro cover cases","16e17e":"funda fundas para cover cases iphone16e iphone17e",
 "17":"funda fundas para forro cover cases iphone17","17 pro":"funda fundas para forro cover cases iphone17pro","17 pro max":"funda fundas para forro cover cases"}
TR="transparent "   # backend word valid for clear series only (Search Volume 554 for 'transparent phone case')
def backend(model,vis_text,clear):
    key=minfo(model)[2]; vt=set(re.findall(r"[a-zñ0-9]+",vis_text.lower()))
    toks=(TR if clear else "").split()+BK[key].split()
    return " ".join(w for w in toks if w not in vt)
def article(w): return "an" if w[0].lower() in "aeiou" else "a"
def remainder(pn,ps):
    full=re.sub(r"\s+"," ",pn).strip()
    if ps.lower()==full.lower(): return ""
    if full.lower().startswith(ps.lower()): rem=full[len(ps):].strip()
    else:
        have={x.lower() for x in ps.split()}; rem=" ".join(w for w in full.split() if w.lower() not in have)
    rem=re.sub(r"[()]","",rem).strip(" ;,")
    return rem[:1].upper()+rem[1:] if rem else ""
def hl_continue(pn,ps,clear,extra_first=()):
    rem=remainder(pn,ps)
    bases=(["Built-in ring compatible with MagSafe chargers and accessories","Ring compatible with MagSafe chargers and accessories","Ring compatible with MagSafe"] if clear
           else ["Built-in magnetic ring compatible with MagSafe chargers and accessories","Magnetic ring compatible with MagSafe chargers","Magnetic ring compatible with MagSafe"])
    extras=list(extra_first)+["2.12 mm TPU back","camera cutout"]
    for base in bases:
        head=(rem+"; " if rem else "")+base
        if len(head)<=125:
            for e in extras:
                if len(head)+2+len(e)<=125: head+="; "+e
            return head
    # remainder too long: shorten it, keep the shortest base
    base=bases[-1]; r=rem
    while r and len(r+"; "+base)>125: r=" ".join(r.split()[:-1])
    return (r+"; " if r else "")+base
def common(model,pn,finish_hl,mt,fits,single):
    hl=f"{pn} print on {article(finish_hl)} {finish_hl} case; fits {mt}; magnetic ring compatible with MagSafe"
    if len(hl)>125: hl=f"{pn} print, {finish_hl} case; fits {mt}; magnet ring compatible with MagSafe"
    if len(hl)>125: hl=f"{pn}, {finish_hl} case; fits {mt}; magnet ring compatible with MagSafe"
    if len(hl)>125:
        tail=f", {finish_hl} case; fits {mt}; magnet ring compatible with MagSafe"
        sp=fit_print(pn,125-len(tail),set()); hl=f"{sp}{tail}"
    return hl
B1="MagSafe compatible magnetic ring: Built-in ring compatible with MagSafe chargers and accessories; N45 magnet, 2300 Gs, 55 mm outer and 46 mm inner diameter"
B5="Slim, lightweight protection: Guards against everyday scratches and bumps while keeping the phone easy to hold"
def build215(sku,model,pn,taken):
    mt,fits,_=minfo(model); single=(fits==mt); pn=clean_pn(pn)
    head=f"{BRAND} Clear Magnetic Case for {mt}"
    ps=fit_print(pn,75-len(head)-2,taken)
    title=f"{head}, {ps}"
    hl=hl_continue(pn,ps,True)
    b=[B1,
       f"Clear {mt} case: Precise fit with a camera cutout that keeps the lens area open" if single else f"Clear {mt} case: Precise fit for {fits} with a camera cutout that keeps the lens area open",
       "Clear TPU back: 2.12 mm thick TPU back panel in a transparent finish, so your phone color shows through the print",
       (f"{pn} design: " if pn[0].isalpha() else "Print design: ")+(f"Decorative print on the clear back, with the phone color visible around the artwork" if pn[0].isalpha() else f"{pn}, a decorative print on the clear back, with the phone color visible around the artwork"),
       B5]
    desc=(f"{BRAND} clear magnetic phone case with a decorative print design for {fits}.\n\n"
          f"Does it work with MagSafe? The case has a built-in magnetic ring compatible with MagSafe chargers and accessories. The ring uses an N45 magnet with 2300 Gs magnetic strength; outer diameter 55 mm, inner diameter 46 mm, thickness 1.4 mm.\n\n"
          f"What is it made of? The back panel is 2.12 mm thick TPU. The camera cutout leaves the lens area open.\n\n"
          f"Is it clear? Yes. The TPU back is clear, so your phone color shows through around the print.\n\n"
          f"What is it for? Everyday protection against scratches and bumps in a slim, lightweight case, and use with MagSafe chargers and accessories.\n\n"
          f"Which phones does it fit? Compatible with {fits}.")
    return us_all(dict(sku=sku,model=model,pn=pn,title=title,hl=hl,bullets=b,desc=desc,clear=True,ps=ps))
F223={"F02":dict(lead="Black Smoky Matte Case",hl="black smoky matte",back="Smoky matte back",txt="2.12 mm thick TPU back panel with a translucent smoky dark gray matte finish, black frame and black buttons",fin="smoky matte back",ask="smoky"),
      "F05":dict(lead="Clear Matte Case",hl="frosted clear matte",back="Frosted clear back",txt="2.12 mm thick TPU back panel with a translucent frosted clear matte finish",fin="frosted clear back",ask="frosted"),
      "F06":dict(lead="Gray Frosted Matte Case",hl="gray frosted matte",back="Frosted gray back",txt="2.12 mm thick TPU back panel with a translucent light frosted gray matte finish, black frame and black buttons",fin="frosted gray back",ask="frosted"),
      "F16":dict(lead="Orange Case",hl="orange soft-touch",back="Solid orange finish",txt="Opaque soft-touch orange case with an orange camera surround and orange buttons; 2.12 mm thick TPU back panel",fin="solid orange back",ask=None)}
def build223(sku,model,pn,fcode,cname,taken):
    f=F223[fcode]; mt,fits,_=minfo(model); single=(fits==mt); pn=clean_pn(pn)
    head=f"{BRAND} {f['lead']} for {mt}"
    ps=fit_print(pn,75-len(head)-2,taken); title=f"{head}, {ps}"
    hl=hl_continue(pn,ps,False,["soft-touch finish"] if fcode=="F16" else [])
    fw="Orange" if fcode=="F16" else "Matte"
    pal=palette(cname)
    b=[B1,
       f"{fw} {mt} case: Precise fit with a camera cutout that keeps the lens area open" if single else f"{fw} {mt} case: Precise fit for {fits} with a camera cutout that keeps the lens area open",
       f"{f['back']}: {f['txt']}",
       (f"{pn} design: " if pn[0].isalpha() else "Print design: ")+(f"Decorative {pal} print on the {f['fin']}" if pn[0].isalpha() else f"{pn}, a decorative {pal} print on the {f['fin']}"),
       B5]
    kind="soft-touch" if fcode=="F16" else "matte"
    finq=("What finish does it have? Solid opaque orange soft-touch finish with an orange camera surround and orange buttons." if fcode=="F16"
          else (f"Is the case matte? Yes. The back is a translucent {f['ask']} matte finish; {f['txt'].split(' finish, ')[1]}." if ' finish, ' in f['txt'] else f"Is the case matte? Yes. The back is a translucent {f['ask']} matte finish."))
    madeof=("Opaque soft-touch case with a 2.12 mm thick TPU back panel" if fcode=="F16" else "2.12 mm thick TPU back panel")
    desc=(f"{BRAND} {f['lead'].lower().replace(' case','')} magnetic phone case with a decorative print design for {fits}.\n\n"
          f"Does it work with MagSafe? The case has a built-in magnetic ring compatible with MagSafe chargers and accessories. The ring uses an N45 magnet with 2300 Gs magnetic strength; outer diameter 55 mm, inner diameter 46 mm, thickness 1.4 mm.\n\n"
          f"{finq}\n\n"
          f"What is it made of? {madeof}. The camera cutout leaves the lens area open.\n\n"
          f"What is it for? Everyday protection against scratches and bumps in a slim, lightweight case, and use with MagSafe chargers and accessories.\n\n"
          f"Which phones does it fit? Compatible with {fits}.")
    return us_all(dict(sku=sku,model=model,pn=pn,title=title,hl=hl,bullets=b,desc=desc,clear=False,ps=ps))
def us_all(o):
    o=dict(o); o['title']=us(o['title']); o['hl']=us(o['hl']); o['bullets']=[us(x) for x in o['bullets']]; o['desc']=us(o['desc']); return o
def checks(o):
    e=[]; t=o['title']
    if len(t)>75: e.append("title>75")
    if re.search(r"[!$?_{}^¬¦]",t): e.append("title special")
    wc=collections.Counter(w for w in re.findall(r"[a-z0-9']+",t.lower()) if w not in{"for","with","and","the","of","a","in"})
    if wc and max(wc.values())>2: e.append("title repeat")
    if len(o['hl'])>125: e.append("hl>125")
    tb=0
    for x in o['bullets']:
        tb+=len(x.encode())
        if not 10<=len(x)<=255: e.append("b len")
        if x.endswith(".") or not x[0].isupper() or re.search(r"[^\x00-\x7f]",x): e.append("b format")
    if tb>1000: e.append("bullets>1000B")
    if len(o['desc'])>2000: e.append("desc>2000")
    return e,tb
def load215():
    names={re.search(r"q(\d+)$",s).group(1):n for s,n in openpyxl.load_workbook(R+"215/Каталог 215 Амазон - SKU prints.xlsx").active.iter_rows(min_row=2,values_only=True)}
    rows=list(openpyxl.load_workbook(R+"215/215 Amazon 29.09 — модели исправлено.xlsx").active.iter_rows(min_row=2,values_only=True))
    pol={r[0]:r for r in openpyxl.load_workbook(R+"215/215 Amazon policy check.xlsx")["Policy check"].iter_rows(min_row=2,values_only=True)}
    out=[]; taken=collections.defaultdict(set)
    for r in rows:
        q=re.search(r"q(\d+)$",r[0]).group(1); o=build215(r[1],r[2],generic_name(names[q]),taken[r[2]]); o['pn_orig']=names[q]
        taken[r[2]].add(o['ps'].lower()); o['q']=q; o['old']=r[0]; o['policy']=pol.get("q"+q); out.append(o)
    return out
def load223():
    w=openpyxl.load_workbook(R+"223/223 Amazon New 17.09 - with prints.xlsx").active
    rows=list(w.iter_rows(min_row=2,values_only=True)); out=[]; seen=set(); taken=collections.defaultdict(set)
    c215={re.search(r"q(\d+)$",s).group(1):n for s,n in openpyxl.load_workbook(R+"215/Каталог 215 Амазон - SKU prints.xlsx").active.iter_rows(min_row=2,values_only=True)}
    pol={r[0]:r for r in openpyxl.load_workbook(R+"215/215 Amazon policy check.xlsx")["Policy check"].iter_rows(min_row=2,values_only=True)}
    for r in rows:
        if r[1] in seen: continue
        seen.add(r[1]); m=re.match(r"223(\w+?)-(F\d+)Ru001q(\d+)",r[0]); f,q=m.group(2),m.group(3)
        key=(r[2],f)
        o=build223(r[1],r[2],generic_name(r[11]),f,r[13],taken[key]); o['pn_orig']=r[11]; taken[key].add(o['ps'].lower()); o['q']=q; o['f']=f; o['old']=r[0]
        o['policy']=pol.get("q"+q) if c215.get(q)==r[11] else None
        out.append(o)
    return out

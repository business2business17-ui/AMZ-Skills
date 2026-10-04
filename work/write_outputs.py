import sys,collections,re; sys.path.insert(0,"/home/user/AMZ-Skills/work")
import gen2,openpyxl
from openpyxl.styles import Font,PatternFill,Alignment
HDR=["SKU","Title","Item Highlight","Bullet 1","Bullet 2","Bullet 3","Bullet 4","Bullet 5","Description","Keywords Backend","Policy risk","Policy type","Policy recommendation (from 215 policy check)"]
RED=PatternFill("solid",fgColor="FFC7CE"); YEL=PatternFill("solid",fgColor="FFEB9C"); BLUE=PatternFill("solid",fgColor="DDEBF7")
def sheet(wb,name,out,clear):
    if name in wb.sheetnames: del wb[name]
    ws=wb.create_sheet(name); ws.append(HDR)
    for o in out:
        vis=" ".join([o['title'],o['hl'],*o['bullets'],o['desc']]); o['backend']=gen2.backend(o['model'],vis,clear)
        p=o.get('policy'); risk,typ,rec=(p[2],p[3],p[5]) if p else ("","","")
        if p and name=="SEO 223": rec="Inherited by same print name from 215 policy check, not checked on 223 artwork. "+str(rec)
        ws.append([o['sku'],o['title'],o['hl'],*o['bullets'],o['desc'],o['backend'],risk,typ,rec]); r=ws.max_row
        if risk=="HIGH":
            for c in ws[r]: c.fill=RED
        elif risk=="MEDIUM": ws.cell(r,11).fill=YEL
    for c in ws[1]: c.font=Font(bold=True); c.fill=BLUE
    ws.freeze_panes="B2"
    for col,w in zip("ABCDEFGHIJKLM",[30,62,45,50,50,50,50,50,70,50,12,24,50]): ws.column_dimensions[col].width=w
def rules(wb,name,rows):
    if name in wb.sheetnames: del wb[name]
    ws=wb.create_sheet(name)
    for r in rows: ws.append(r)
    for c in ws[1]: c.font=Font(bold=True); c.fill=BLUE
    for col,w in zip("ABC",[26,110,62]): ws.column_dimensions[col].width=w
    for row in ws.iter_rows(min_row=2):
        for c in row: c.alignment=Alignment(wrap_text=True,vertical="top")
COMMON=[("Title (max 75)","Brand + product type + color/finish + model + print name; print name is cut at a clean word boundary when it does not fit; no ! $ ? _ { } ^ ¬ ¦; no word more than twice","CONFIRMED: Amazon title requirements text (user-provided md, also in repo SEO-IPhone-Case)"),
("Item Highlight (max 125)","[Print name] print on a [finish] case; fits [model]; magnetic ring compatible with MagSafe","Limit user-provided"),
("Bullets (5)","Header: description; capital first letter; no end period; no emoji or special characters; 10-255 chars each; first 1000 bytes of all five indexed; no repeats between bullets","Format CONFIRMED (Amazon text); 255 and 1000-byte limits user-provided"),
("Description (max 2000)","Plain short paragraphs with question-answer lines for Rufus/COSMO readability; facts only","2000 user-provided; Rufus/COSMO/A9-A10 reliance is industry inference (COSMO itself is a real Amazon system, SIGMOD 2024)"),
("Backend (max 250 bytes)","Only words from phrases with Search Volume >= 500 (Helium 10, 2026-10-03) that are not already in visible text","Helium 10 analyze_keywords"),
("Print names","Taken from the owner's catalog tables built from the PDF catalogs (old SKU = catalog SKU, model code replaced by xxx). Old Excel Color / Pattern names and the old per-print description sentences are NOT used (138 of 149 names differed from catalog, descriptions not verified against pictures). Names spot-checked against catalog pictures (215 page 1, 223 pages 1 and 12): match","Owner rule 2026-10-04: print must come from the PDF catalogs"),
("MagSafe wording","magnetic ring compatible with MagSafe chargers and accessories; never certified, never Made for MagSafe, never Apple","Owner decision"),
("Facts used","TPU back 2.12 mm; camera cutout; magnetic ring N45, 2300 Gs, outer 55 mm, inner 46 mm, thickness 1.4 mm; slim and lightweight (owner confirmed)","Owner TTX"),
("Claims not used","shockproof, air-cushion corners, bumper, raised edges, waterproof, rust-proof, military grade, certified","Not confirmed by owner TTX"),
("Model fit","iPhone 18 Pro = 17 Pro, 18 Pro Max = 17 Pro Max, 16e = 17e (identical dimensions)","Owner statement 2026-10-04, not independently verified"),
("Policy risk columns","Columns K-M carry the risk level from '215 Amazon policy check.xlsx' (HIGH rows are filled red: the owner's own recommendation is do not upload until the artwork is changed)","Owner policy check")]
R=gen2.R
out=gen2.load215(); wb=openpyxl.load_workbook(R+"215/215 Amazon 29.09 — модели исправлено.xlsx"); sheet(wb,"SEO 215",out,True)
rules(wb,"SEO Rules 215",[("Field","Rule applied in series 215 (clear MagSafe cases)","Source / status")]+COMMON+[
("Models","11 models as in '215 Amazon 29.09 - модели исправлено.xlsx' (1639 SKU). iPhone 16 and 16 Pro Max (298 SKU) were removed from the table by the owner and are not written","Owner file"),
("Duplicates","Catalog shows q177/q199 and q189/q202/q207 are different prints (sizes), so no duplicate-title workaround is used; every title is unique","Catalog print names")])
wb.save("/home/user/AMZ-Skills/work/215_Amazon_29.09_SEO.xlsx")
out2=gen2.load223()
UNCHK={"002","003","006","008","029","056","078","171","172","173","174","175","176","197"}
NOTES={"078":"Recognizable painting (looks like American Gothic by Grant Wood): Claude observation, verify rights before upload",
       "171":"Park name in artwork text (Acadia): Claude observation, verify trademark/licensing","172":"Park name in artwork text (Grand Canyon): Claude observation, verify trademark/licensing",
       "173":"Park name in artwork text (Joshua Tree): Claude observation, verify trademark/licensing","174":"Park name in artwork text (Yellowstone): Claude observation, verify trademark/licensing",
       "175":"Park name in artwork text (Yosemite): Claude observation, verify trademark/licensing","176":"Park name in artwork text (Zion): Claude observation, verify trademark/licensing"}
for o in out2:
    if o.get("policy") is None and o["q"] in UNCHK:
        o["policy"]=(o["q"],o["pn"],"CHECK","Not covered by 215 policy check","",NOTES.get(o["q"],"Print exists only in series 223, not covered by the 215 policy check"))
wb=openpyxl.load_workbook(R+"223/223 Amazon New 17.09 - with prints.xlsx"); sheet(wb,"SEO 223",out2,False)
rules(wb,"SEO Rules 223",[("Field","Rule applied in series 223 (colored MagSafe cases)","Source / status")]+COMMON+[
("Color / finish in title","F02 = Black Smoky Matte (black frame), F05 = Frosted Clear Matte (frame color NOT stated: catalog shows white on some prints and taupe on others), F06 = Gray Frosted Matte (black frame), F16 = Orange (no Matte; solid soft-touch). Follows the owner's Case Color column (Black / Clear / Grey (Black Frame) / Orange)","Owner file 'with prints' (Case Color), catalog pictures"),
("Print colors","Bullet 4 and description use the Color Name column (e.g. Black / Red) as the print palette","Owner file"),
("Matte / frosted / translucent","Used on purpose in title, Item Highlight, bullets 2-3 and description Q&A although Cerebro has no such case phrases and Helium 10 shows SV >= 500 only for the 17 line","Owner decision 2026-10-04"),
("Duplicates","13 exact duplicate rows of the source table removed (2739 -> 2726 SKU)","Source table")])
wb.save("/home/user/AMZ-Skills/work/223_Amazon_with_prints_SEO.xlsx")
print("saved",len(out),len(out2))
